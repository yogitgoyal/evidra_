import csv
import io
from datetime import date, datetime
from typing import Any

from openpyxl import load_workbook


def _stringify(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    return str(value).strip()


def _normalize_rows(
    headers: list[str],
    rows: list[tuple[int, list[Any]]],
) -> list[tuple[int, dict[str, str]]]:
    normalized_headers = [_stringify(header) for header in headers]
    if not any(normalized_headers):
        raise ValueError("CSV/XLSX file must include a header row.")

    result = []
    for line_number, values in rows:
        row = {
            header: _stringify(values[index]) if index < len(values) else ""
            for index, header in enumerate(normalized_headers)
            if header
        }
        if any(row.values()):
            result.append((line_number, row))
    return result


def _extract_csv(content: bytes) -> list[tuple[int, dict[str, str]]]:
    try:
        text = content.decode("utf-8-sig")
    except UnicodeDecodeError as error:
        raise ValueError("CSV file must be UTF-8 encoded.") from error

    reader = csv.reader(io.StringIO(text))
    try:
        headers = next(reader)
    except StopIteration as error:
        raise ValueError("CSV file must include a header row.") from error
    return _normalize_rows(headers, list(enumerate(reader, start=2)))


def _extract_xlsx(content: bytes) -> list[tuple[int, dict[str, str]]]:
    workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    try:
        worksheet = workbook.active
        rows = worksheet.iter_rows(values_only=True)
        try:
            headers = list(next(rows))
        except StopIteration as error:
            raise ValueError("XLSX file must include a header row.") from error
        return _normalize_rows(headers, list(enumerate(rows, start=2)))
    finally:
        workbook.close()


def extract_cdr_rows(
    filename: str,
    content_type: str | None,
    content: bytes,
) -> list[tuple[int, dict[str, str]]]:
    lowered_name = filename.casefold()
    if content_type == "text/csv" or lowered_name.endswith(".csv"):
        return _extract_csv(content)
    if (
        content_type
        == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        or lowered_name.endswith(".xlsx")
    ):
        return _extract_xlsx(content)
    raise ValueError("Only CSV and XLSX files are supported.")
